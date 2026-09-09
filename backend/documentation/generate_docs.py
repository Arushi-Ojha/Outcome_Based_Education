import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
import os

def create_api_docs():
    doc = docx.Document()
    
    # Title
    title = doc.add_heading('OBE Tracking System - API & Pipeline Documentation', 0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    
    # 1. Workflow Pipeline
    doc.add_heading('1. Workflow Pipeline', level=1)
    pipeline_text = (
        "The OBE (Outcome Based Education) Tracking System operates on a strict hierarchical workflow to ensure "
        "data integrity and appropriate authorization at every step.\n\n"
        "Phase 1: Institution & Role Setup (Admin & HOD)\n"
        "- The HOD signs up and creates a Department.\n"
        "- Faculties and Coordinators sign up, and wait in a 'PENDING' state.\n"
        "- The HOD approves them, which automatically generates their respective profiles (Faculties table).\n\n"
        "Phase 2: Curriculum & OBE Framework Blueprinting (HOD)\n"
        "- The HOD defines the Academic Courses (Programs).\n"
        "- The HOD creates Subjects and maps them to semesters (Mandatory & Elective Baskets).\n"
        "- The HOD defines GAs, PEOs, POs, and KSA Templates (with exact Knowledge, Skill, Attitude weightages).\n"
        "- The HOD maps PEOs to GAs, and POs to PEOs.\n\n"
        "Phase 3: Course Blueprinting (Coordinator)\n"
        "- Coordinators are assigned to Subjects by the HOD.\n"
        "- The Coordinator defines Course Outcomes (COs) and maps them to POs.\n"
        "- The Coordinator defines Assessments and maps them to the COs.\n"
        "- The Coordinator assigns Faculties to teach specific Subjects.\n\n"
        "Phase 4: Execution (Faculty)\n"
        "- Assigned Faculties bulk upload Student Details for a semester. The engine automatically maps students to the Faculty's subjects.\n"
        "- Faculties bulk upload IA (Internal Assessment) Marks for those students.\n\n"
        "Phase 5: Attainment Calculation (Analytics)\n"
        "- The system calculates attainment bottom-up: Assessments -> COs -> POs -> PEOs -> GAs."
    )
    doc.add_paragraph(pipeline_text)
    
    # 2. Attainment Calculations
    doc.add_heading('2. Attainment Calculations Engine', level=1)
    calc_text = (
        "The system uses a weighted, bottom-up approach to calculate attainment across the OBE framework. "
        "The calculations are performed automatically using the defined thresholds and mapping weightages.\n\n"
        "1. Assessment -> CO Attainment\n"
        "For each CO, the system aggregates the marks from all Assessments mapped to that CO.\n"
        "- A student 'attains' the CO if their total score crosses the Assessment's threshold_percentage (e.g., 65%).\n"
        "- Overall CO Attainment % = (Number of students passing the threshold / Total enrolled students) * 100.\n\n"
        "2. CO -> PO/PSO Attainment\n"
        "A PO's attainment is derived from all COs mapped to it.\n"
        "- PO Attainment = Sum(CO_Attainment * Mapping_Weight) / Sum(Mapping_Weights)\n"
        "- If a CO is tagged with a KSA Template, the attainment is further split into Knowledge, Skill, and Attitude "
        "buckets based on the floats defined in the KSA_Tag (e.g., K_weight: 4.5, S_weight: 0.5).\n\n"
        "3. PO -> PEO Attainment\n"
        "PEO Attainment = Sum(PO_Attainment * Mapping_Weight) / Sum(Mapping_Weights)\n\n"
        "4. PEO -> GA Attainment\n"
        "GA Attainment = Sum(PEO_Attainment * Mapping_Weight) / Sum(Mapping_Weights)"
    )
    doc.add_paragraph(calc_text)
    
    # 3. API Endpoints
    doc.add_heading('3. API Endpoints Reference', level=1)
    
    def add_endpoint(method, url, desc):
        p = doc.add_paragraph()
        run_method = p.add_run(f"[{method}] ")
        run_method.bold = True
        run_method.font.color.rgb = RGBColor(0, 102, 204)
        run_url = p.add_run(url)
        run_url.bold = True
        p.add_run(f"\n{desc}")
        
    doc.add_heading('Auth Domain', level=2)
    add_endpoint("POST", "/auth/request-otp", "Requests a signup OTP to the provided email.")
    add_endpoint("POST", "/auth/signup", "Registers a new user (HOD, Faculty, Coordinator).")
    add_endpoint("POST", "/auth/login", "Authenticates the user and returns a JWT token.")
    
    doc.add_heading('HOD Domain', level=2)
    add_endpoint("POST", "/hod/departments", "Creates a department and links the HOD to it.")
    add_endpoint("GET", "/hod/staff/pending", "Lists all staff members awaiting approval in the department.")
    add_endpoint("PATCH", "/hod/staff/{user_id}/approve", "Approves a staff member and auto-generates their Faculty profile.")
    add_endpoint("POST", "/hod/academic-courses", "Defines a new academic course (program).")
    add_endpoint("POST", "/hod/subjects", "Creates a new subject in the department.")
    add_endpoint("POST", "/hod/academic-courses/mapping", "Bulk maps subjects to academic courses (Curriculum Blueprint).")
    add_endpoint("POST", "/hod/obe/gas", "Defines Graduate Attributes (GAs).")
    add_endpoint("POST", "/hod/obe/peos", "Defines Program Educational Objectives (PEOs).")
    add_endpoint("POST", "/hod/obe/pos", "Defines Program Outcomes (POs/PSOs) with a broad KSA Domain.")
    add_endpoint("POST", "/hod/obe/ksas", "Defines detailed KSA Tag templates (with Knowledge, Skill, Attitude weightages).")
    add_endpoint("POST", "/hod/obe/mapping/peo-ga", "Bulk maps PEOs to GAs.")
    add_endpoint("POST", "/hod/obe/mapping/po-peo", "Bulk maps POs to PEOs.")
    
    doc.add_heading('Coordinator Domain', level=2)
    add_endpoint("POST", "/coordinator/course-outcomes", "Defines Course Outcomes (COs) mapped to specific KSA Tags.")
    add_endpoint("POST", "/coordinator/mapping/co-po", "Bulk maps COs to POs.")
    add_endpoint("POST", "/coordinator/assessments", "Defines an assessment structure for a subject.")
    add_endpoint("POST", "/coordinator/mapping/assessment-co", "Maps specific assessments to COs.")
    add_endpoint("POST", "/coordinator/subjects/{subject_id}/assign-faculty/{user_id}", "Assigns a Faculty member to teach a Subject.")
    
    doc.add_heading('Faculty Domain', level=2)
    add_endpoint("POST", "/faculty/students/bulk-upload", "Registers and auto-enrolls students into subjects matching their semester and the Faculty's assignments.")
    add_endpoint("POST", "/faculty/subjects/{subject_id}/marks", "Bulk uploads IA marks for students against specific assessments.")
    
    doc.add_heading('System Health & Integration', level=2)
    add_endpoint("GET", "/", "API root greeting and verification endpoint.")
    add_endpoint("GET", "/api/health", "System health check returning operational status for frontend and monitoring.")

    # 4. Frontend & CORS Integration
    doc.add_heading('4. Frontend & CORS Integration', level=1)
    frontend_text = (
        "The React JavaScript frontend is located in the 'frontend' directory, built with Vite and linted using Oxlint.\n\n"
        "1. CORS Architecture:\n"
        "- Configured via FastAPI CORSMiddleware in 'backend/main/main.py'.\n"
        "- Whitelists dev origins ('http://localhost:5173', 'http://127.0.0.1:5173', 'http://localhost:3000', 'http://127.0.0.1:3000').\n"
        "- Uses regex '^https?://(localhost|127\\.0\\.0\\.1)(:\\d+)?$' to authorize dynamic localhost ports.\n"
        "- Supports credentials, all HTTP methods, and custom Authorization/Content-Type headers.\n\n"
        "2. Frontend Client & Proxy:\n"
        "- Vite Dev Server proxy maps '/api' to 'http://127.0.0.1:8000'.\n"
        "- Resilient API helper ('frontend/src/api/client.js') handles health checks and connection validation.\n"
        "- Linted with Oxlint ('npm run lint') with zero errors and zero warnings."
    )
    doc.add_paragraph(frontend_text)

    # Save Document
    file_path = os.path.join(os.getcwd(), 'OBE_API_Documentation.docx')
    doc.save(file_path)
    print(f"Documentation generated successfully at: {file_path}")


if __name__ == "__main__":
    create_api_docs()
